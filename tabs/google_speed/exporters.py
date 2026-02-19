import io
from typing import Dict, Iterable, List

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


DEFAULT_CATEGORY_ORDER = [
    "performance",
    "accessibility",
    "best-practices",
    "seo",
]

CATEGORY_LABELS = {
    "performance": "Performance",
    "accessibility": "Accessibility",
    "best-practices": "Best Practices",
    "seo": "SEO",
}


def _ordered_strategies(strategies: List[str]) -> List[str]:
    ordered = []
    for strategy in ("mobile", "desktop"):
        if strategy in strategies:
            ordered.append(strategy)
    return ordered


def _ordered_categories(categories: List[str]) -> List[str]:
    ordered = []
    for category in DEFAULT_CATEGORY_ORDER:
        if category in categories:
            ordered.append(category)
    return ordered


def _format_score(row: Dict, category: str):
    scores = row.get("scores", {}) or {}
    score = scores.get(category)
    if row.get("status") == "error":
        return "ERROR"
    if score is None:
        return "N/A"
    return score


def rows_to_pagespeed_xlsx_bytes(
    rows: Iterable[dict],
    strategies: List[str],
    categories: List[str],
) -> bytes:
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl не установлена. Установите: pip install openpyxl")

    strategy_list = _ordered_strategies(strategies)
    if not strategy_list:
        strategy_list = ["mobile"]

    category_list = _ordered_categories(categories)
    if not category_list:
        category_list = ["performance"]

    rows_list = list(rows)

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for strategy in strategy_list:
        ws = wb.create_sheet(title="Mobile" if strategy == "mobile" else "Desktop")
        headers = ["URL", "Status", "Error"] + [
            CATEGORY_LABELS.get(cat, cat.title()) for cat in category_list
        ]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        current_row = 2
        for row in rows_list:
            if row.get("strategy") != strategy:
                continue

            values = [
                row.get("url", ""),
                row.get("status", ""),
                row.get("error", ""),
            ]
            for category in category_list:
                values.append(_format_score(row, category))

            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 60
        for index in range(4, len(headers) + 1):
            col = chr(64 + index)
            ws.column_dimensions[col].width = 16
        ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
